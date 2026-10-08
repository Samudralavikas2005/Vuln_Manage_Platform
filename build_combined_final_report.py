import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import re
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:left w:val="none"/>\n'
            f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:right w:val="none"/>\n'
            f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:insideV w:val="none"/>\n'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_formatted_runs(paragraph, text, font_size=10.5, font_name='Arial', default_color=(0x33, 0x41, 0x55)):
    """
    Parses inline markdown like **bold**, *italic*, and `code` into docx runs.
    """
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)')
    tokens = pattern.split(text)
    
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            r = paragraph.add_run(token[2:-2])
            r.bold = True
            r.font.name = font_name
            r.font.size = Pt(font_size)
            r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        elif token.startswith('*') and token.endswith('*'):
            r = paragraph.add_run(token[1:-1])
            r.italic = True
            r.font.name = font_name
            r.font.size = Pt(font_size)
            r.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        elif token.startswith('`') and token.endswith('`'):
            r = paragraph.add_run(token[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(font_size - 1)
            r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        else:
            r = paragraph.add_run(token)
            r.font.name = font_name
            r.font.size = Pt(font_size)
            r.font.color.rgb = RGBColor(*default_color)

def parse_markdown_to_docx(stage_files, output_path):
    doc = docx.Document()
    
    # 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # Title Banner
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(6)
    r_t = p_title.add_run("Software Security Development Lifecycle (SSDLC)")
    r_t.font.name = 'Arial'
    r_t.font.size = Pt(24)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(16)
    r_s = p_sub.add_run("Comprehensive End-to-End Final Report (Stages 1 – 4)\nVulnShield Enterprise Vulnerability Management Platform")
    r_s.font.name = 'Arial'
    r_s.font.size = Pt(13)
    r_s.font.italic = True
    r_s.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(36)
    r_m = p_meta.add_run("Prepared by: Enterprise Security Engineering Team\nDate: October 2026  |  Status: Consolidated Final Report")
    r_m.font.size = Pt(9.5)
    r_m.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    doc.add_page_break()

    for f_idx, filepath in enumerate(stage_files):
        if not os.path.exists(filepath):
            print(f"Warning: File {filepath} not found. Skipping.")
            continue
            
        with open(filepath, 'r', encoding='utf-8') as f:
            lines = f.readlines()
            
        if f_idx > 0:
            doc.add_page_break()

        in_code_block = False
        code_lines = []
        code_lang = ""
        
        in_table = False
        table_rows = []

        i = 0
        while i < len(lines):
            line = lines[i].rstrip('\n')
            
            # --- CODE BLOCK HANDLING ---
            if line.strip().startswith('```'):
                if not in_code_block:
                    in_code_block = True
                    code_lang = line.strip()[3:].strip()
                    code_lines = []
                else:
                    in_code_block = False
                    # Render code block table
                    t = doc.add_table(rows=1, cols=1)
                    t.alignment = WD_TABLE_ALIGNMENT.CENTER
                    cell = t.rows[0].cells[0]
                    set_cell_background(cell, "F1F5F9")
                    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
                    cell.width = Inches(6.5)
                    
                    set_table_borders(t, color="CBD5E1", sz="4")
                    
                    p_code = cell.paragraphs[0]
                    p_code.paragraph_format.space_after = Pt(0)
                    p_code.paragraph_format.line_spacing = 1.1
                    
                    code_text = "\n".join(code_lines)
                    r_code = p_code.add_run(code_text)
                    r_code.font.name = 'Consolas'
                    r_code.font.size = Pt(8.5)
                    r_code.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                    
                    # Space after code block
                    p_spacer = doc.add_paragraph()
                    p_spacer.paragraph_format.space_after = Pt(6)
                i += 1
                continue

            if in_code_block:
                code_lines.append(line)
                i += 1
                continue

            # --- TABLE HANDLING ---
            if '|' in line and line.strip().startswith('|') and line.strip().endswith('|'):
                if not in_table:
                    in_table = True
                    table_rows = []
                # Check if it's a separator line like |---|---|
                if re.match(r'^\|[\s\-:|]+\|$', line.strip()):
                    i += 1
                    continue
                
                # Split cells
                raw_cells = [c.strip() for c in line.strip().split('|')[1:-1]]
                table_rows.append(raw_cells)
                i += 1
                continue
            else:
                if in_table:
                    # Flush table to document
                    in_table = False
                    if table_rows:
                        num_cols = max(len(r) for r in table_rows)
                        t = doc.add_table(rows=1, cols=num_cols)
                        t.alignment = WD_TABLE_ALIGNMENT.CENTER
                        set_table_borders(t, color="CBD5E1", sz="4")
                        
                        # Header
                        hdr_cells = t.rows[0].cells
                        hdr_data = table_rows[0]
                        for c_idx in range(num_cols):
                            cell_txt = hdr_data[c_idx] if c_idx < len(hdr_data) else ""
                            hdr_cells[c_idx].text = cell_txt
                            set_cell_background(hdr_cells[c_idx], "0F172A")
                            set_cell_margins(hdr_cells[c_idx], top=120, bottom=120, left=120, right=120)
                            p = hdr_cells[c_idx].paragraphs[0]
                            for run in p.runs:
                                run.font.name = 'Arial'
                                run.font.bold = True
                                run.font.size = Pt(9.0)
                                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                                
                        # Rows
                        for r_idx, r_data in enumerate(table_rows[1:]):
                            row_cells = t.add_row().cells
                            bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                            for c_idx in range(num_cols):
                                cell_txt = r_data[c_idx] if c_idx < len(r_data) else ""
                                row_cells[c_idx].text = ""
                                p_c = row_cells[c_idx].paragraphs[0]
                                p_c.paragraph_format.space_after = Pt(2)
                                p_c.paragraph_format.line_spacing = 1.15
                                add_formatted_runs(p_c, cell_txt, font_size=8.5, font_name='Arial')
                                set_cell_background(row_cells[c_idx], bg)
                                set_cell_margins(row_cells[c_idx], top=90, bottom=90, left=120, right=120)
                                
                        p_sp = doc.add_paragraph()
                        p_sp.paragraph_format.space_after = Pt(8)
                    table_rows = []

            # --- EMPTY LINE ---
            if not line.strip():
                i += 1
                continue

            # --- HORIZONTAL RULE ---
            if line.strip() in ['---', '***', '___']:
                p_hr = doc.add_paragraph()
                p_hr.paragraph_format.space_before = Pt(8)
                p_hr.paragraph_format.space_after = Pt(8)
                pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E1"/></w:pBdr>')
                p_hr._element.get_or_add_pPr().append(pBdr)
                i += 1
                continue

            # --- HEADINGS ---
            if line.startswith('# '):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(20)
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.keep_with_next = True
                add_formatted_runs(p, line[2:].strip(), font_size=16, default_color=(0x0F, 0x17, 0x2A))
                for r in p.runs:
                    r.bold = True
                i += 1
                continue
            elif line.startswith('## '):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(16)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.keep_with_next = True
                add_formatted_runs(p, line[3:].strip(), font_size=13.5, default_color=(0x1E, 0x3A, 0x8A))
                for r in p.runs:
                    r.bold = True
                i += 1
                continue
            elif line.startswith('### '):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.keep_with_next = True
                add_formatted_runs(p, line[4:].strip(), font_size=11.5, default_color=(0x33, 0x41, 0x55))
                for r in p.runs:
                    r.bold = True
                i += 1
                continue
            elif line.startswith('#### '):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.keep_with_next = True
                add_formatted_runs(p, line[5:].strip(), font_size=10.5, default_color=(0x47, 0x55, 0x69))
                for r in p.runs:
                    r.bold = True
                i += 1
                continue

            # --- BULLETS AND NUMBERS ---
            bullet_match = re.match(r'^([\*\-\+])\s+(.*)$', line.strip())
            number_match = re.match(r'^(\d+\.)\s+(.*)$', line.strip())
            
            if bullet_match:
                p = doc.add_paragraph(style='List Bullet')
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.left_indent = Inches(0.25)
                add_formatted_runs(p, bullet_match.group(2))
                i += 1
                continue
            elif number_match:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.left_indent = Inches(0.25)
                r_num = p.add_run(number_match.group(1) + " ")
                r_num.bold = True
                r_num.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                add_formatted_runs(p, number_match.group(2))
                i += 1
                continue

            # --- CALLOUT / BLOCKQUOTE ---
            if line.strip().startswith('>'):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.right_indent = Inches(0.3)
                pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="12" w:color="1E3A8A"/></w:pBdr>')
                p._element.get_or_add_pPr().append(pBdr)
                add_formatted_runs(p, line.strip()[1:].strip(), default_color=(0x47, 0x55, 0x69))
                i += 1
                continue

            # --- STANDARD PARAGRAPH ---
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_runs(p, line.strip())
            i += 1

        # Check if table ended at end of file
        if in_table and table_rows:
            num_cols = max(len(r) for r in table_rows)
            t = doc.add_table(rows=1, cols=num_cols)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(t, color="CBD5E1", sz="4")
            
            hdr_cells = t.rows[0].cells
            hdr_data = table_rows[0]
            for c_idx in range(num_cols):
                cell_txt = hdr_data[c_idx] if c_idx < len(hdr_data) else ""
                hdr_cells[c_idx].text = cell_txt
                set_cell_background(hdr_cells[c_idx], "0F172A")
                set_cell_margins(hdr_cells[c_idx], top=120, bottom=120, left=120, right=120)
                p = hdr_cells[c_idx].paragraphs[0]
                for run in p.runs:
                    run.font.name = 'Arial'
                    run.font.bold = True
                    run.font.size = Pt(9.0)
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    
            for r_idx, r_data in enumerate(table_rows[1:]):
                row_cells = t.add_row().cells
                bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                for c_idx in range(num_cols):
                    cell_txt = r_data[c_idx] if c_idx < len(r_data) else ""
                    row_cells[c_idx].text = ""
                    p_c = row_cells[c_idx].paragraphs[0]
                    p_c.paragraph_format.space_after = Pt(2)
                    p_c.paragraph_format.line_spacing = 1.15
                    add_formatted_runs(p_c, cell_txt, font_size=8.5, font_name='Arial')
                    set_cell_background(row_cells[c_idx], bg)
                    set_cell_margins(row_cells[c_idx], top=90, bottom=90, left=120, right=120)

    doc.save(output_path)
    print(f"[+] Final Consolidated Word Document successfully compiled at: {output_path}")

if __name__ == "__main__":
    docs_dir = "/home/vikas/SSE_END_LAB/docs"
    stages = [
        os.path.join(docs_dir, "STAGE_1_DESIGN_AND_ARCHITECTURE.md"),
        os.path.join(docs_dir, "STAGE_2_THREAT_AND_AGILE.md"),
        os.path.join(docs_dir, "STAGE_3_CODE_AND_CONTAINERIZATION.md"),
        os.path.join(docs_dir, "STAGE_4_TESTING_AND_FINAL_REVIEW.md")
    ]
    out_file = os.path.join(docs_dir, "SSDLC_Final_Comprehensive_Report_Stages_1_to_4.docx")
    parse_markdown_to_docx(stages, out_file)
